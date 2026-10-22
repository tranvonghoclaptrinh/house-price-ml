"""Đo lại cấu hình GỐC, chỉ xuất bảng mới; không lưu/ghi đè model.

Chạy ở thư mục gốc: python notebooks_or_src_py/chuan_hoa_ket_qua.py
"""
from pathlib import Path
import hashlib
import json
import os
import platform
import subprocess
import time
from datetime import datetime, timezone

import joblib
import numpy as np
import pandas as pd
import sklearn
import xgboost
from sklearn.model_selection import cross_val_score
from sklearn.pipeline import Pipeline
from metrics import evaluate_all
from preprocess import load_and_split, build_pipeline, RANDOM_SEED
from train_compare import get_models
from tien_te import TY_GIA_USD_VND, sang_vnd, dinh_dang_vnd, ghi_chu_quy_doi

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/chuan_hoa'
ORDER = ['XGBoost', 'Gradient Boosting', 'Random Forest', 'Linear Regression', 'Decision Tree']
ALLOWED = {'service/predict_service.py', 'public/index.html', 'requirements.txt', 'README.md'}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snapshot():
    """Ghi nguồn một lần; chạy lại không được hợp thức hóa thay đổi tệp gốc."""
    manifest = OUT / 'nguon_goc_sha256.json'
    if manifest.exists():
        saved = json.loads(manifest.read_text(encoding='utf-8'))
        for name, digest in saved['files'].items():
            if name not in ALLOWED:
                assert sha(ROOT / name) == digest, f'Tệp gốc đã thay đổi: {name}'
        return saved
    names = subprocess.check_output(['git', 'ls-files', '-z'], cwd=ROOT).decode().split('\0')
    # On a fresh clone this is the original inventory. Never include our new files on a rerun.
    names = [n for n in names if n and not n.startswith('results/chuan_hoa/')
             and n not in {'notebooks_or_src_py/tien_te.py', 'notebooks_or_src_py/chuan_hoa_ket_qua.py',
                           'notebooks_or_src_py/verify_chuan_hoa.py', 'notebooks_or_src_py/du_doan_vnd.py'}
             and not n.endswith('.bak')]
    saved = {'base_commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT).decode().strip(),
             'files': {n: sha(ROOT / n) for n in names}}
    manifest.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding='utf-8')
    return saved


def check_originals(saved):
    for name, digest in saved['files'].items():
        if name not in ALLOWED:
            assert sha(ROOT / name) == digest, f'Không được thay đổi: {name}'


def metrics(y, pred):
    values = evaluate_all(y, pred)
    values['MAPE'] = float(np.mean(np.abs((np.asarray(y, float)-np.asarray(pred, float))/np.asarray(y, float)))*100)
    return values


def gap_label(ratio, gap):
    if ratio >= 2.0 or gap >= .10: return 'Khoảng cách lớn'
    if ratio >= 1.5 or gap >= .06: return 'Khoảng cách vừa'
    if ratio >= 1.25 or gap >= .03: return 'Khoảng cách nhỏ'
    return 'Chưa thấy dấu hiệu rõ'


def save_csv(name, df):
    df.to_csv(OUT / name, index=False, encoding='utf-8-sig')
    return df


def export_original_vnd():
    source = pd.read_csv(ROOT / 'results/model_results.csv')
    source = source.set_index('Model').loc[ORDER].reset_index()
    result = pd.DataFrame({'Mo_hinh': source.Model, 'MAE_VND': sang_vnd(source.MAE),
                          'RMSE_VND': sang_vnd(source.RMSE), 'R2': source.R2,
                          'CV_MAE_mean_VND': sang_vnd(source.CV_MAE_mean),
                          'Train_time_s': source.Train_time_s,
                          'Tap_danh_gia': 'TEST (MAE/RMSE/R2); CV trên TRAIN (CV_MAE_mean_VND)',
                          'Don_vi_sai_so': 'VND'})
    return save_csv('model_results_VND.csv', result)


def write_workbook(tables, wide, readme):
    # openpyxl explicitly permitted by the project specification; runs on Windows as well.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    cover = wb.active; cover.title = 'Đọc trước tiên'
    for line in readme.splitlines(): cover.append([line.lstrip("# ").replace("`", "")])
    cover.column_dimensions['A'].width = 115
    cover.freeze_panes = 'A2'
    for row in cover:
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
        row[0].font = Font(name='Calibri', size=11, bold=str(row[0].value or '').startswith('#'))
        cover.row_dimensions[row[0].row].height = 32 if len(str(row[0].value or '')) > 110 else 22
    labels = {'Id': 'Id', 'Mo_hinh': 'Mô hình', 'Tap': 'Tập dữ liệu', 'So_dong': 'Số dòng',
              'MAE_VND': 'MAE (₫)', 'RMSE_VND': 'RMSE (₫)', 'R2': 'R²',
              'MAPE_phan_tram': 'MAPE (%)', 'MAE_trieu_VND': 'MAE (triệu ₫)',
              'SalePrice_USD_goc': 'SalePrice USD gốc', 'SalePrice_VND': 'SalePrice (₫)',
              'Thuc_te_VND': 'Thực tế (₫)', 'Du_doan_VND': 'Dự đoán (₫)',
              'Sai_so_VND': 'Sai số (₫)', 'Sai_so_tuyet_doi_VND': 'Sai số tuyệt đối (₫)',
              'Sai_so_phan_tram': 'Sai số (%)', 'Trung_binh_VND': 'CV MAE trung bình (₫)',
              'Do_lech_chuan_VND': 'CV MAE độ lệch chuẩn (₫)', 'MAE_TRAIN_VND': 'MAE TRAIN (₫)',
              'MAE_TEST_VND': 'MAE TEST (₫)', 'Ti_le_TEST_chia_TRAIN': 'Tỷ lệ TEST / TRAIN',
              'R2_TRAIN': 'R² TRAIN', 'R2_TEST': 'R² TEST', 'Chenh_R2': 'Chênh R²',
              'Nhan_khoang_cach': 'Khoảng cách tham khảo', 'Tham_so': 'Toàn bộ tham số gốc',
              'Ngay_chay': 'Ngày chạy UTC', 'So_fold_CV': 'Số fold CV',
              'Tap_danh_gia': 'Phạm vi đánh giá', 'Don_vi_sai_so': 'Đơn vị sai số',
              'Nhom': 'Nhóm dữ liệu', 'CV_MAE_mean_VND': 'CV trên TRAIN MAE (₫)',
              'Train_time_s': 'Thời gian fit gốc (giây)', 'count': 'Số căn', 'mean': 'Giá trung bình (₫)', 'std': 'Độ lệch chuẩn giá (₫)', 'min': 'Giá nhỏ nhất (₫)', '25%': 'Phân vị 25% (₫)', '50%': 'Trung vị giá (₫)', '75%': 'Phân vị 75% (₫)', 'max': 'Giá lớn nhất (₫)'}
    entries = list(tables.items()) + [('Dự đoán TEST (dạng rộng)', wide)]
    for sheet_name, df in entries:
        ws = wb.create_sheet(sheet_name)
        headers = [labels.get(str(c), str(c).replace('_', ' ')) for c in df.columns]
        ws.append(headers)
        for record in df.itertuples(index=False, name=None):
            ws.append([v.item() if isinstance(v, np.generic) else v for v in record])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 36
        for cell in ws[1]:
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='284B63')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for j, col in enumerate(df.columns, 1):
            is_percent = col in ('MAPE_phan_tram', 'Sai_so_phan_tram')
            is_money = ('VND' in str(col) and 'trieu' not in str(col)) or str(col).endswith('(₫)') or (sheet_name == 'Thống kê giá' and col not in ('Nhom', 'count'))
            fmt = '0.00%' if is_percent else ('0.0000' if 'R2' in str(col) else '#,##0' if is_money else '0.00' if pd.api.types.is_float_dtype(df[col]) else 'General')
            for cells in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                c = cells[0]
                if is_percent and isinstance(c.value, (int, float)): c.value /= 100
                c.number_format = fmt
                c.font = Font(name='Calibri', size=11)
                c.alignment = Alignment(vertical='top', wrap_text=col in ('Tham_so','Tap_danh_gia'))
            longest = max([len(headers[j-1])] + [len(str(x)) for x in df[col].head(200)])
            ws.column_dimensions[get_column_letter(j)].width = min(45, max(18, longest+2))
        # Required currency note, separate from the filterable data region.
        line = len(df)+3
        ws.cell(line, 1, ghi_chu_quy_doi())
        ws.merge_cells(start_row=line, start_column=1, end_row=line, end_column=max(2, len(df.columns)))
        ws.cell(line,1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[line].height = 38
        if sheet_name == 'Cấu hình':
            import math
            ws.column_dimensions['B'].width = 95
            for k in range(2, len(df)+2):
                ws.row_dimensions[k].height = min(409, max(100, math.ceil(len(str(ws.cell(k,2).value))/75)*15+30))
        if 'Tap_danh_gia' in df.columns:
            for k in range(2, len(df)+2): ws.row_dimensions[k].height = 48
        if sheet_name == 'Thống kê giá':
            for cell in ws['B'][1:len(df)+1]: cell.number_format = '#,##0'
    import io
    import zipfile
    buffer = io.BytesIO()
    wb.save(buffer)
    payload = buffer.getvalue()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.testzip() is None, 'Workbook có phần ZIP hỏng'
    temporary = OUT / '.workbook.tmp'
    with temporary.open('wb') as stream:
        stream.write(payload); stream.flush(); os.fsync(stream.fileno())
    temporary.replace(OUT / 'BAO_CAO_KET_QUA_TRAIN_TEST.xlsx')


def plots(train, test, detail, prices):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.family'] = 'DejaVu Sans'
    loc = np.arange(len(ORDER)); a = train.MAE_VND.to_numpy()/1e6; b = test.MAE_VND.to_numpy()/1e6
    fig, ax = plt.subplots(figsize=(12, 6)); width = .36
    for values, offset, label, color in [(a,-width/2,'TRAIN','#6B98A8'),(b,width/2,'TEST','#D78744')]:
        bars = ax.bar(loc+offset,values,width,label=label,color=color)
        ax.bar_label(bars, labels=[f'{v:.1f}'.replace('.',',') for v in values],padding=3)
    ax.set_xticks(loc, ORDER); ax.set_ylabel('MAE (triệu đồng quy đổi)'); ax.set_title('Sai số MAE trên TRAIN và TEST của năm mô hình gốc')
    ax.legend(); ax.set_ylim(0,max(b)*1.16); fig.tight_layout(); fig.savefig(OUT/'bieu_do_MAE_TRAIN_vs_TEST_VND.png',dpi=160); plt.close(fig)
    x = detail[detail.Mo_hinh == 'XGBoost']; actual=x.Thuc_te_VND.to_numpy()/1e9; pred=x.Du_doan_VND.to_numpy()/1e9
    fig,ax=plt.subplots(figsize=(7,7)); ax.scatter(actual,pred,alpha=.65,s=20,color='#356480'); low=min(actual.min(),pred.min()); high=max(actual.max(),pred.max());ax.plot([low,high],[low,high],'--',color='#BB543A',label='Dự đoán bằng thực tế')
    ax.set(xlabel='Giá nhãn TEST (tỷ đồng quy đổi)',ylabel='Giá dự đoán (tỷ đồng quy đổi)',title='XGBoost gốc trên TEST');ax.legend();fig.tight_layout();fig.savefig(OUT/'bieu_do_thuc_te_vs_du_doan_TEST_XGBoost.png',dpi=160);plt.close(fig)
    fig,ax=plt.subplots(figsize=(10,5));ax.hist(sang_vnd(prices)/1e9,bins=35,color='#467C91',edgecolor='white');ax.set(xlabel='SalePrice (tỷ đồng quy đổi)',ylabel='Số căn',title='Phân phối SalePrice toàn bộ dữ liệu có nhãn');fig.tight_layout();fig.savefig(OUT/'bieu_do_phan_phoi_SalePrice_VND.png',dpi=160);plt.close(fig)


def main():
    os.chdir(ROOT); OUT.mkdir(parents=True, exist_ok=True)
    saved = snapshot(); start = time.perf_counter()
    Xtr,Xte,ytr,yte=load_and_split(); originals=pd.read_csv('results/model_results.csv').set_index('Model')
    models=get_models(); assert set(models)==set(ORDER), 'Cần đúng năm mô hình từ get_models()'
    preprocessor=build_pipeline(Xtr)
    records={'TRAIN':[], 'TEST':[]}; details={'TRAIN':[], 'TEST':[]}; cv_rows=[]; config=[]; evidence={}; now=datetime.now(timezone.utc).isoformat()
    stored=joblib.load('models/best_model.pkl'); stored_predictions=np.asarray(stored.predict(Xte),float)
    for name in ORDER:
        print(f'Đo lại cấu hình gốc và CV: {name}', flush=True)
        model=models[name]; pipe=Pipeline(steps=[('preprocess',preprocessor),('model',model)])
        pipe.fit(Xtr,ytr)
        predictions={scope:np.asarray(pipe.predict(X),float) for scope,X in [('TRAIN',Xtr),('TEST',Xte)]}
        usd={}
        for scope,X,y in [('TRAIN',Xtr,ytr),('TEST',Xte,yte)]:
            pred=predictions[scope]; usd[scope]=metrics(y,pred); val=usd[scope]
            records[scope].append({'Mo_hinh':name,'Tap':scope,'So_dong':len(y),'MAE_VND':sang_vnd(val['MAE']),
                                  'RMSE_VND':sang_vnd(val['RMSE']),'R2':val['R2'],'MAPE_phan_tram':val['MAPE'],'MAE_trieu_VND':sang_vnd(val['MAE'])/1e6})
            for source_id,actual,guess in zip(X.Id,y,pred):
                e=sang_vnd(float(guess)-float(actual))
                details[scope].append({'Id':int(source_id),'Mo_hinh':name,'Thuc_te_VND':sang_vnd(float(actual)),
                                       'Du_doan_VND':sang_vnd(float(guess)),'Sai_so_VND':e,'Sai_so_tuyet_doi_VND':abs(e),'Sai_so_phan_tram':abs(e)/sang_vnd(float(actual))*100})
        for key,digits in [('MAE',2),('RMSE',2),('R2',4)]:
            assert round(usd['TEST'][key],digits)==originals.loc[name,key], f'{name}: {key} khác kết quả gốc'
        fold=-cross_val_score(pipe,Xtr,ytr,cv=5,scoring='neg_mean_absolute_error')
        assert round(float(fold.mean()),2)==originals.loc[name,'CV_MAE_mean'], f'{name}: CV khác kết quả gốc'
        cv_rows.append({'Mo_hinh':name,'Tap':'CV trên TRAIN',**{f'Fold_{i+1}_VND':sang_vnd(float(v)) for i,v in enumerate(fold)},'Trung_binh_VND':sang_vnd(float(fold.mean())),'Do_lech_chuan_VND':sang_vnd(float(fold.std()))})
        usd['CV_folds']=fold.tolist(); evidence[name]=usd
        config.append({'Mo_hinh':name,'Tham_so':json.dumps(model.get_params(),ensure_ascii=False,default=str,sort_keys=True),
                       'random_state':model.get_params().get('random_state', 'Không có'), 'test_size':len(Xte)/(len(Xtr)+len(Xte)),
                       'So_fold_CV':len(fold),'scikit-learn':sklearn.__version__,'xgboost':xgboost.__version__,'pandas':pd.__version__,'numpy':np.__version__,'Ngay_chay':now})
        if name=='XGBoost':
            err=np.abs(predictions['TEST']-stored_predictions)
            assert np.allclose(predictions['TEST'],stored_predictions,rtol=1e-6,atol=0), 'XGBoost khác artifact đã lưu'
            evidence['_xgboost_compare']={'max_abs_USD_goc':float(err.max()),'max_relative':float(np.max(err/np.maximum(abs(stored_predictions),np.finfo(float).tiny))),'predictions_USD_goc':stored_predictions.tolist()}
    df=pd.read_csv('data/train.csv'); membership=pd.DataFrame({'Id':df.Id,'Tap':np.where(df.index.isin(Xtr.index),'TRAIN','TEST'),'SalePrice_USD_goc':df.SalePrice,'SalePrice_VND':sang_vnd(df.SalePrice)})
    train=pd.DataFrame(records['TRAIN']);test=pd.DataFrame(records['TEST']);cv=pd.DataFrame(cv_rows)
    gap=pd.DataFrame([{'Mo_hinh':n,'MAE_TRAIN_VND':sang_vnd(evidence[n]['TRAIN']['MAE']),'MAE_TEST_VND':sang_vnd(evidence[n]['TEST']['MAE']),
                      'Ti_le_TEST_chia_TRAIN':evidence[n]['TEST']['MAE']/evidence[n]['TRAIN']['MAE'],'R2_TRAIN':evidence[n]['TRAIN']['R2'],'R2_TEST':evidence[n]['TEST']['R2'],
                      'Chenh_R2':evidence[n]['TRAIN']['R2']-evidence[n]['TEST']['R2'],
                      'Nhan_khoang_cach':gap_label(evidence[n]['TEST']['MAE']/evidence[n]['TRAIN']['MAE'],evidence[n]['TRAIN']['R2']-evidence[n]['TEST']['R2'])} for n in ORDER])
    stats=pd.DataFrame([{'Nhom':scope,**sang_vnd(series).describe().to_dict()} for scope,series in [('TOÀN BỘ',df.SalePrice),('TRAIN',ytr),('TEST',yte)]])
    frames=[membership,train,test,cv,gap,pd.DataFrame(details['TEST']),pd.DataFrame(details['TRAIN']),pd.DataFrame(config),stats]
    names=['01_phan_chia_du_lieu.csv','02_ket_qua_TRAIN.csv','03_ket_qua_TEST.csv','04_ket_qua_CV_5fold.csv','05_so_sanh_TRAIN_vs_TEST.csv','06_du_doan_TEST_chi_tiet.csv','07_du_doan_TRAIN_chi_tiet.csv','08_cau_hinh_mo_hinh.csv','09_thong_ke_SalePrice_VND.csv']
    for filename,frame in zip(names,frames):save_csv(filename,frame)
    converted=export_original_vnd()
    note=ghi_chu_quy_doi()
    readme='# Kết quả TRAIN TEST và quy đổi VND\n\n'+note+'\n\n'
    descriptions=['Danh sách Id và giá nhãn gốc, cho biết từng căn thuộc TRAIN hay TEST.','Sai số trên TRAIN đã dùng học; không dùng làm bằng chứng chất lượng tổng quát.','Sai số trên TEST: các dòng không nằm trong phần fit của lần chia này.','CV 5 fold trên TRAIN, cv=5 không shuffle, scoring=neg_mean_absolute_error như repo gốc.','So sánh riêng TRAIN và TEST để nhận diện khoảng cách, ngưỡng chỉ tham khảo.','Từng dự đoán TEST: sai số bằng dự đoán trừ nhãn.','Từng dự đoán TRAIN, có thể lạc quan hơn TEST.','Tham số nhập từ get_models() nguyên trạng, phiên bản và ngày đo lại.','Thống kê SalePrice theo VND; count là số căn, không nhân tỷ giá.']
    for filename,desc in zip(names,descriptions): readme+=f'- `{filename}`: {desc}\n'
    row=test.iloc[0];readme+=f'\nVí dụ đọc một dòng: {row.Mo_hinh} – TEST: sai số tuyệt đối trung bình {dinh_dang_vnd(row.MAE_VND)} trên {int(row.So_dong)} căn không thuộc phần fit của lần chia này.\n\n'
    readme+='`model_results_VND.csv` quy đổi các số đã làm tròn trong CSV gốc. File 02/03 và chi tiết giữ số tính lại đầy đủ; chênh vài chục đồng là do CSV gốc làm tròn USD đến 2 chữ số. Không dùng số quy đổi từ CSV làm tròn thay số chi tiết.\n\n'
    readme+='Mô hình giữ nguyên được huấn luyện trên USD gốc. Nhân giá bằng một hằng số không làm đổi R² và MAPE; MAE và RMSE chỉ đổi đơn vị. MSE nếu xem riêng phải nhân bình phương tỷ giá và ghi VND².\n\n'
    readme+=f'Chia dữ liệu: test_size={len(Xte)/(len(Xtr)+len(Xte))}; random_state={RANDOM_SEED}. Giữ nguyên Id trong đặc trưng theo yêu cầu bảo toàn bản gốc. Không đổi cấu hình, không tuning, không lưu .pkl.\n\n'
    readme+='Ngưỡng mang tính tham khảo: lớn nếu TEST/TRAIN MAE ≥ 2,0 hoặc chênh R² ≥ 0,10; vừa nếu ≥ 1,5 hoặc ≥ 0,06; nhỏ nếu ≥ 1,25 hoặc ≥ 0,03; còn lại chưa thấy dấu hiệu rõ. Ưu tiên mức lớn trước.\n\n'
    readme+='Workbook gom các CSV vào sheet riêng, tiền hiển thị nguyên đồng, R² 4 chữ số; MAPE/sai số phần trăm lưu dạng tỷ lệ trong Excel để định dạng phần trăm đúng. CSV giữ tên cột máy đọc theo đặc tả; Excel có tiêu đề tiếng Việt. Ba PNG mới thể hiện MAE TRAIN/TEST, XGBoost TEST và phân phối giá, không thay ảnh cũ.\n\n'
    readme+='Hạn chế giữ nguyên: bản gốc chọn best_model bằng MAE TEST; không gọi TEST là đánh giá cuối chưa từng dùng lựa chọn. Id vẫn ảnh hưởng đầu vào. Nhãn là SalePrice trong CSV, chưa xác minh thị trường bên ngoài. Thời gian ở model_results_VND giữ từ CSV gốc, không phải thời gian chạy máy hiện tại.\n'
    (OUT/'00_DOC_TRUOC_TIEN.md').write_text(readme,encoding='utf-8')
    wide=frames[5].pivot(index='Id',columns='Mo_hinh',values='Du_doan_VND').reindex(columns=ORDER)
    wide.insert(0,'Thực tế (₫)',frames[5].drop_duplicates('Id').set_index('Id').Thuc_te_VND)
    wide=wide.rename(columns={n:n+' (₫)' for n in ORDER}).reset_index()
    sheet_names=['Phân chia dữ liệu','Kết quả TRAIN','Kết quả TEST','CV trên TRAIN','So sánh TRAIN TEST','Dự đoán TEST chi tiết','Dự đoán TRAIN chi tiết','Cấu hình','Thống kê giá']
    write_workbook(dict(zip(sheet_names,frames)) | {'Quy đổi CSV gốc':converted},wide,readme)
    plots(train,test,frames[5],df.SalePrice)
    evidence['_meta']={'rate':TY_GIA_USD_VND,'seed':RANDOM_SEED,'generated_at':now,'elapsed_seconds':time.perf_counter()-start,'base_commit':saved['base_commit'],'train_indices':Xtr.index.tolist(),'test_indices':Xte.index.tolist()}
    (OUT/'so_lieu_goc_va_doi_chieu.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    check_originals(saved)
    print(test[['Mo_hinh','MAE_VND','R2']].to_string(index=False)); print(f'Hoàn tất: {evidence["_meta"]["elapsed_seconds"]:.2f} giây; không ghi model.')


if __name__ == '__main__': main()
