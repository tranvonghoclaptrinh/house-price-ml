"""Kiểm chứng độc lập các bảng chuẩn hóa; luôn ghi PASS/FAIL và số thật."""
from pathlib import Path
import hashlib
import json
import os
import sys
from datetime import datetime, timezone

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.model_selection import train_test_split
from tien_te import TY_GIA_USD_VND, sang_vnd, ghi_chu_quy_doi

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/chuan_hoa'
ORDER=['XGBoost','Gradient Boosting','Random Forest','Linear Regression','Decision Tree']
# References are acceptance tolerances from the supplied spec, never output metrics.
TRAIN_MAE_REFERENCE=dict(zip(ORDER,[6087,9667,6479,15467,10639]))
TRAIN_R2_REFERENCE=dict(zip(ORDER,[.9881,.9707,.9798,.9071,.9647]))
TEST_REFERENCE=dict(zip(ORDER,[(15822.88,26086.62,.9113,16908.23),(17099.28,26755.67,.9067,16638.37),(17634.87,29183.87,.8890,18301.01),(20594.94,31143.09,.8736,20188.39),(25754.33,40546.96,.7857,26322.78)]))


def read(name):return pd.read_csv(OUT/name,encoding='utf-8-sig',float_precision='round_trip')
def close(a,b,rtol=1e-9,atol=1e-6):return bool(np.allclose(a,b,rtol=rtol,atol=atol))


def main():
    os.chdir(ROOT); OUT.mkdir(parents=True,exist_ok=True); log=[]
    def check(title, ok, actual):
        log.append(('PASS' if bool(ok) else 'FAIL',title,str(actual)))
    try:
        raw=pd.read_csv(ROOT/'data/train.csv'); original=pd.read_csv(ROOT/'results/model_results.csv').set_index('Model')
        split=read('01_phan_chia_du_lieu.csv'); train=read('02_ket_qua_TRAIN.csv').set_index('Mo_hinh'); test=read('03_ket_qua_TEST.csv').set_index('Mo_hinh');cv=read('04_ket_qua_CV_5fold.csv').set_index('Mo_hinh')
        evidence=json.loads((OUT/'so_lieu_goc_va_doi_chieu.json').read_text(encoding='utf-8'))
        seed=evidence['_meta']['seed'];tr,te=train_test_split(raw,test_size=.2,random_state=seed)
        ids_train=set(split.loc[split.Tap=='TRAIN','Id']);ids_test=set(split.loc[split.Tap=='TEST','Id'])
        check('Kích thước và phân hoạch TRAIN/TEST',len(ids_train)==1168 and len(ids_test)==292 and len(split)==len(raw)==1460 and not ids_train&ids_test and ids_train|ids_test==set(raw.Id),f'TRAIN={len(ids_train)}, TEST={len(ids_test)}, tổng={len(split)}, giao={len(ids_train&ids_test)}')
        check('Id TEST tái lập từ split gốc',ids_test==set(te.Id),f'seed={seed}; sai khác đối xứng={len(ids_test ^ set(te.Id))}')
        check('Giá nhãn và quy đổi trong phân hoạch',close(split.SalePrice_USD_goc,raw.SalePrice,rtol=0,atol=0) and close(split.SalePrice_VND,sang_vnd(raw.SalePrice),rtol=0,atol=0),f'sai khác VND tối đa={float(np.max(abs(split.SalePrice_VND-sang_vnd(raw.SalePrice))))}')
        for name in ORDER:
            src=original.loc[name]; rr=test.loc[name]; folds=cv.loc[name,[f'Fold_{i}_VND' for i in range(1,6)]].to_numpy(dtype=float)
            check(name+' TEST khớp CSV gốc',abs(rr.MAE_VND/TY_GIA_USD_VND-src.MAE)<=.01 and abs(rr.RMSE_VND/TY_GIA_USD_VND-src.RMSE)<=.01 and round(rr.R2,4)==src.R2,f'MAE USD={rr.MAE_VND/TY_GIA_USD_VND:.12f}; gốc={src.MAE}; RMSE USD={rr.RMSE_VND/TY_GIA_USD_VND:.12f}; gốc={src.RMSE}; R²={rr.R2:.12f}; gốc={src.R2}')
            ref=TEST_REFERENCE[name]
            check(name+' TEST khớp tham chiếu đặc tả',abs(rr.MAE_VND/TY_GIA_USD_VND-ref[0])<=.01 and abs(rr.RMSE_VND/TY_GIA_USD_VND-ref[1])<=.01 and round(rr.R2,4)==ref[2] and abs(cv.loc[name,'Trung_binh_VND']/TY_GIA_USD_VND-ref[3])<=.01,f'MAE={rr.MAE_VND/TY_GIA_USD_VND}; RMSE={rr.RMSE_VND/TY_GIA_USD_VND}; R²={rr.R2}; CV={cv.loc[name,"Trung_binh_VND"]/TY_GIA_USD_VND}')
            check(name+' CV 5 fold trên TRAIN',close(folds.mean(),cv.loc[name,'Trung_binh_VND']) and close(folds.std(),cv.loc[name,'Do_lech_chuan_VND']) and abs(folds.mean()/TY_GIA_USD_VND-src.CV_MAE_mean)<=.01,f'fold VND={folds.tolist()}; mean USD={folds.mean()/TY_GIA_USD_VND}; gốc={src.CV_MAE_mean}')
            for scope,summary,filename,ids in [('TEST',test,'06_du_doan_TEST_chi_tiet.csv',ids_test),('TRAIN',train,'07_du_doan_TRAIN_chi_tiet.csv',ids_train)]:
                detail=read(filename);part=detail[detail.Mo_hinh==name]; actual=part.Thuc_te_VND;guess=part.Du_doan_VND;err=guess-actual
                mae=float(err.abs().mean());rmse=float(np.sqrt((err**2).mean()));r2=float(1-(err**2).sum()/((actual-actual.mean())**2).sum());mape=float((err.abs()/actual).mean()*100)
                row=summary.loc[name];usd=evidence[name][scope]
                check(name+' '+scope+' tính lại độc lập',len(part)==len(ids) and set(part.Id)==ids and part.Id.is_unique and close(mae,row.MAE_VND) and close(rmse,row.RMSE_VND) and close(r2,row.R2) and close(mape,row.MAPE_phan_tram),f'n={len(part)}; MAE={mae}; RMSE={rmse}; R²={r2}; MAPE={mape}')
                check(name+' '+scope+' từng sai số đúng',close(err,part.Sai_so_VND) and close(err.abs(),part.Sai_so_tuyet_doi_VND) and close(err.abs()/actual*100,part.Sai_so_phan_tram),f'sai khác cột sai số tối đa={float(np.max(abs(err-part.Sai_so_VND)))} VND')
                check(name+' '+scope+' đổi đơn vị đúng',close(row.MAE_VND,sang_vnd(usd['MAE']),atol=0) and close(row.RMSE_VND,sang_vnd(usd['RMSE']),atol=0) and close(row.R2,usd['R2'],atol=0) and close(row.MAPE_phan_tram,usd['MAPE'],atol=0),f'MAE VND={row.MAE_VND}; MAE USD gốc={usd["MAE"]}; R² giữ nguyên={row.R2}; MAPE giữ nguyên={row.MAPE_phan_tram}')
            trrow=train.loc[name]
            check(name+' TRAIN đối chiếu tham chiếu',trrow.MAE_VND>0 and abs(trrow.MAE_VND/TY_GIA_USD_VND-TRAIN_MAE_REFERENCE[name])<=1 and abs(trrow.R2-TRAIN_R2_REFERENCE[name])<=.0005,f'MAE TRAIN USD={trrow.MAE_VND/TY_GIA_USD_VND}; R² TRAIN={trrow.R2}')
        stored=joblib.load(ROOT/'models/best_model.pkl');pred=np.asarray(stored.predict(te.drop(columns='SalePrice')),float)
        detail=read('06_du_doan_TEST_chi_tiet.csv');p=detail[detail.Mo_hinh=='XGBoost'].set_index('Id').loc[te.Id,'Du_doan_VND'].to_numpy()/TY_GIA_USD_VND
        diff=abs(pred-p);check('XGBoost đo lại trùng best_model.pkl',close(pred,p,rtol=1e-6,atol=0),f'sai khác USD tối đa={diff.max()}; sai khác tương đối tối đa={np.max(diff/np.maximum(abs(pred),np.finfo(float).tiny))}; lúc đo trước xuất CSV={evidence["_xgboost_compare"]["max_abs_USD_goc"]}')
        for name,digest in json.loads((OUT/'nguon_goc_sha256.json').read_text(encoding='utf-8'))['files'].items():
            actual=hashlib.sha256((ROOT/name).read_bytes()).hexdigest()
            if actual==digest:check('SHA-256 '+name,True,actual)
            elif name in {'service/predict_service.py','public/index.html','requirements.txt','README.md'}:
                bak=ROOT/(name+'.bak');ok=bak.is_file() and hashlib.sha256(bak.read_bytes()).hexdigest()==digest
                if name in {'requirements.txt','README.md'}:ok=ok and (ROOT/name).read_bytes().startswith(bak.read_bytes())
                check('Thêm có bản gốc '+name,ok,f'gốc={digest}; mới={actual}; backup khớp={ok}')
            else:check('SHA-256 '+name,False,f'gốc={digest}; mới={actual}')
        files=sorted(OUT.glob('*.csv'))
        for file in files:
            frame=read(file.name);check('CSV UTF-8 BOM '+file.name,file.read_bytes().startswith(b'\xef\xbb\xbf') and '\ufffd' not in file.read_text(encoding='utf-8-sig'),f'{len(frame)} dòng; {len(frame.columns)} cột; BOM={file.read_bytes().startswith(bytes([239,187,191]))}')
            bad=[str(c) for c in frame.columns if 'USD' in str(c) and not (str(c).endswith('_USD_goc') or str(c).endswith('_goc'))]
            check('Đơn vị cột chính '+file.name,not bad,f'cột USD không ghi gốc={bad}')
            if 'Mo_hinh' in frame:check('Thứ tự model '+file.name,list(frame.Mo_hinh.drop_duplicates())==ORDER,str(list(frame.Mo_hinh.drop_duplicates())))
        wb=load_workbook(OUT/'BAO_CAO_KET_QUA_TRAIN_TEST.xlsx',data_only=False)
        check('Workbook mở được và có sheet bắt buộc',len(wb.sheetnames)==12 and wb.sheetnames[0]=='Đọc trước tiên' and 'Dự đoán TEST (dạng rộng)' in wb.sheetnames,str(wb.sheetnames))
        for ws in wb:
            values=[c.value for row in ws for c in row if c.value is not None]
            check('Excel ghi chú và ký tự '+ws.title,any(ghi_chu_quy_doi() in str(v) for v in values) and not any('\ufffd' in str(v) for v in values),f'hàng={ws.max_row}; cột={ws.max_column}; freeze={ws.freeze_panes}')
        ws=wb['Kết quả TEST'];check('Định dạng tiền R² phần trăm Excel',ws['D2'].number_format=='#,##0' and ws['F2'].number_format=='0.0000' and ws['G2'].number_format=='0.00%' and close(ws['G2'].value*100,test.iloc[0].MAPE_phan_tram),f'D2={ws["D2"].number_format}; F2={ws["F2"].number_format}; G2={ws["G2"].number_format}; MAPE lưu tỷ lệ={ws["G2"].value}')
        wide=wb['Dự đoán TEST (dạng rộng)'];check('Dự đoán TEST dạng rộng',wide.max_row==len(te)+3 and wide.max_column==7,f'hàng={wide.max_row} gồm header và ghi chú; cột={wide.max_column}')
        converted=read('model_results_VND.csv').set_index('Mo_hinh')
        for n in ORDER:
            a=converted.loc[n];b=original.loc[n];check('Quy đổi CSV gốc '+n,close(a.MAE_VND,sang_vnd(b.MAE),atol=0) and close(a.RMSE_VND,sang_vnd(b.RMSE),atol=0) and close(a.CV_MAE_mean_VND,sang_vnd(b.CV_MAE_mean),atol=0) and a.R2==b.R2 and a.Train_time_s==b.Train_time_s,f'MAE={a.MAE_VND}; RMSE={a.RMSE_VND}; CV={a.CV_MAE_mean_VND}; R²={a.R2}; thời gian gốc={a.Train_time_s}')
        gap=read('05_so_sanh_TRAIN_vs_TEST.csv').set_index('Mo_hinh')
        for n in ORDER:
            a=gap.loc[n]; ratio=test.loc[n,'MAE_VND']/train.loc[n,'MAE_VND'];delta=train.loc[n,'R2']-test.loc[n,'R2']
            label='Khoảng cách lớn' if ratio>=2 or delta>=.10 else 'Khoảng cách vừa' if ratio>=1.5 or delta>=.06 else 'Khoảng cách nhỏ' if ratio>=1.25 or delta>=.03 else 'Chưa thấy dấu hiệu rõ'
            check('Khoảng cách '+n,close(a.Ti_le_TEST_chia_TRAIN,ratio) and close(a.Chenh_R2,delta) and a.Nhan_khoang_cach==label,f'tỷ lệ={ratio}; gap R²={delta}; nhãn={a.Nhan_khoang_cach}')
        stats=read('09_thong_ke_SalePrice_VND.csv')
        for _,row in stats.iterrows():
            series=raw.SalePrice if row.Nhom=='TOÀN BỘ' else tr.SalePrice if row.Nhom=='TRAIN' else te.SalePrice
            expected=sang_vnd(series).describe();check('Thống kê giá '+row.Nhom,all(close(row[k],v) for k,v in expected.items()),f'count={row["count"]}; mean VND={row["mean"]}; min={row["min"]}; max={row["max"]}')
        for file in OUT.glob('bieu_do_*.png'):
            from PIL import Image
            with Image.open(file) as im:im.verify()
            check('Biểu đồ '+file.name,True,f'{file.stat().st_size} bytes, PNG đọc được')
    except Exception as exc:
        check('Không có lỗi khi kiểm chứng',False,f'{type(exc).__name__}: {exc}')
    count=sum(s=='PASS' for s,_,_ in log);failed=len(log)-count
    text=f'# Kiểm chứng kết quả TRAIN TEST và VND\n\n{ghi_chu_quy_doi()}\n\nNgày kiểm tra UTC: {datetime.now(timezone.utc).isoformat()}\n\nKết quả: {count} PASS; {failed} FAIL.\n\n'
    text+='\n'.join(f'- {status} — {title}: {value}' for status,title,value in log)
    text+='\n\nCác phép đo chỉ tái lập mô hình gốc. Không bỏ Id, không thay pipeline, không tuning; TEST từng được bản gốc dùng chọn model, nên không coi là đánh giá cuối chưa từng dùng lựa chọn.\n'
    (OUT/'KIEM_CHUNG.md').write_text(text,encoding='utf-8')
    print(f'{count} PASS; {failed} FAIL → results/chuan_hoa/KIEM_CHUNG.md')
    for status,title,value in log:
        if status=='FAIL': print(f'{status}: {title}: {value}')
    return 1 if failed else 0


if __name__=='__main__':sys.exit(main())
